# -*- coding: utf-8 -*-
"""
Actualiza el Chartpack de Expectativas (estilo Estudios Económicos Santander)
con datos de la BDE del Banco Central de Chile y del GSCPI de la Fed de Nueva York.

Uso:
    export BCCH_USER="correo@registrado.cl"
    export BCCH_PASS="api key de la BDE"
    python actualizar_chartpack.py                 # genera PPTX (y PDF si hay LibreOffice)
    python actualizar_chartpack.py --buscar "imce precio"   # busca códigos en el catálogo

Los códigos de series se resuelven contra el catálogo de la BDE (SearchSeries).
Cada serie puede fijarse con `codigo=...` en SERIES; si no, se busca por palabras clave.
El detalle de lo resuelto queda en salida/series_resueltas.csv para revisión.
"""

import argparse
import datetime as dt
import io
import json
import os
import shutil
import subprocess
import sys
import time
import unicodedata

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import pandas as pd
import requests
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(DIR, "assets")
CACHE = os.path.join(DIR, "cache")
SALIDA = os.path.join(DIR, "salida")

BDE_URL = "https://si3.bcentral.cl/SieteRestWS/SieteRestWS.ashx"
GSCPI_URL = ("https://www.newyorkfed.org/medialibrary/research/interactives/"
             "gscpi/downloads/gscpi_data.xlsx")

###############################################################################
# Estilo Santander (colores tomados del chartpack original)                   #
###############################################################################
ROJO_TITULO = "#EC0000"
AZUL = "#0E2841"
ROJO = "#FF0000"
ROJO_OSC = "#C00000"
VERDE = "#00B050"
AZUL_CLARO = "#0070C0"
DORADO = "#E1AF19"
PETROLEO = "#156082"
NARANJO = "#E97132"
CELESTE = "#0F9ED5"
VERDE_SUAVE = "#64B464"
MORADO = "#A02B93"
GRIS = "#595959"
GRIS_LINEA = "#BFBFBF"

FUENTE = "Calibri"
MESES = ["ene", "feb", "mar", "abr", "may", "jun",
         "jul", "ago", "sep", "oct", "nov", "dic"]
MESES_LARGO = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
               "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def configurar_fuente():
    # Calibri si está instalada; Carlito es métricamente equivalente.
    disponibles = {f.name for f in font_manager.fontManager.ttflist}
    for nombre in ("Calibri", "Carlito"):
        if nombre in disponibles:
            plt.rcParams["font.family"] = nombre
            break
    plt.rcParams.update({
        "font.size": 9,
        "axes.edgecolor": GRIS_LINEA,
        "axes.linewidth": 0.8,
        "xtick.color": GRIS,
        "ytick.color": GRIS,
        "xtick.labelsize": 9,
        "ytick.labelsize": 9,
        "legend.fontsize": 9,
        "legend.frameon": False,
        "axes.spines.top": False,
        "axes.spines.right": False,
    })


###############################################################################
# Series                                                                      #
###############################################################################
# palabras: todas deben aparecer en el título en español (sin tildes, minúsculas);
# "a|b" significa alternativa. codigo: fija la serie y omite la búsqueda.
SERIES = {
    # IMCE
    "imce_pre_com": dict(frec=["MONTHLY"], palabras=["imce", "precio", "comercio"], excluir=["sin miner"]),
    "imce_pre_con": dict(frec=["MONTHLY"], palabras=["imce", "precio", "construcci"]),
    "imce_pre_ind": dict(frec=["MONTHLY"], palabras=["imce", "precio", "industria"]),
    "imce_cos_com": dict(frec=["MONTHLY"], palabras=["imce", "costo", "comercio"]),
    "imce_cos_con": dict(frec=["MONTHLY"], palabras=["imce", "costo", "construcci"]),
    "imce_cos_ind": dict(frec=["MONTHLY"], palabras=["imce", "costo", "industria"]),
    "imce_inv": dict(frec=["MONTHLY"], palabras=["imce", "inversi"], excluir=["sin miner", "comercio", "construcci", "industria", "mineria"]),
    # Expectativas de inflación
    "eee_1": dict(frec=["MONTHLY"], palabras=["expectativ|eee", "inflaci|ipc", "11 meses"]),
    "eee_2": dict(frec=["MONTHLY"], palabras=["expectativ|eee", "inflaci|ipc", "23 meses"]),
    "eof_1": dict(frec=["DAILY", "MONTHLY"], palabras=["eof|operadores financieros", "inflaci|ipc", "11 meses|1 ano"]),
    "eof_2": dict(frec=["DAILY", "MONTHLY"], palabras=["eof|operadores financieros", "inflaci|ipc", "23 meses|2 anos"]),
    # Compensación inflacionaria
    "comp_1": dict(frec=["DAILY"], palabras=["compensaci", "inflaci", "1 ano"], excluir=["swap", "1 en 1"]),
    "comp_2": dict(frec=["DAILY"], palabras=["compensaci", "inflaci", "2 anos"], excluir=["swap", "1 en 1"]),
    "comp_1en1_swap": dict(frec=["DAILY"], palabras=["compensaci|inflaci", "swap", "1 en 1"]),
    # EDEP: costos de la empresa
    "edep_cos_lab": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "costos laborales"]),
    "edep_cos_ins": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "precio|insumo", "insumos"]),
    "edep_cos_maq": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "maquinaria"]),
    "edep_cos_fij": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "costos fijos"]),
    # EDEP: influencia sobre el precio de venta
    "edep_inf_ven": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "nivel|niveles", "venta"]),
    "edep_inf_dol": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "dolar|tipo de cambio"]),
    "edep_inf_mar": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "margen|margenes"]),
    "edep_inf_dis": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "disponibilidad"]),
    "edep_inf_cos": dict(frec=["MONTHLY"], palabras=["edep|determinantes y expectativas de precios", "influencia|incidencia", "costo"]),
    # EPN: evolución de costos (últimos 3 meses)
    "epn_cos_lab": dict(frec=["QUARTERLY", "MONTHLY"], palabras=["epn|percepciones de negocios", "costos laborales"]),
    "epn_cos_ins": dict(frec=["QUARTERLY", "MONTHLY"], palabras=["epn|percepciones de negocios", "insumos"]),
    "epn_cos_maq": dict(frec=["QUARTERLY", "MONTHLY"], palabras=["epn|percepciones de negocios", "maquinaria"]),
    "epn_cos_fij": dict(frec=["QUARTERLY", "MONTHLY"], palabras=["epn|percepciones de negocios", "costos fijos"]),
    # EDEP: expectativas de inflación (códigos publicados por el BCCh)
    "edep_inf_1": dict(frec=["MONTHLY"], codigo="F089.EDEP.POR.91EI12.M"),
    "edep_inf_2": dict(frec=["MONTHLY"], codigo="F089.EDEP.POR.91EI24.M"),
    # Tipo de cambio
    "tcn": dict(frec=["DAILY"], codigo="F073.TCO.PRE.Z.D"),
    "tcm": dict(frec=["DAILY"], palabras=["tipo de cambio multilateral"], excluir=["tcm-5", "tcm-x", "sin", "real"]),
    # Tasas swap promedio cámara
    "spc_clp_1": dict(frec=["DAILY"], palabras=["swap", "camara|spc", "pesos|nominal", "1 ano|360"]),
    "spc_clp_2": dict(frec=["DAILY"], palabras=["swap", "camara|spc", "pesos|nominal", "2 anos|720"]),
    "spc_uf_1": dict(frec=["DAILY"], palabras=["swap", "camara|spc", "uf|reajustable", "1 ano|360"]),
    "spc_uf_2": dict(frec=["DAILY"], palabras=["swap", "camara|spc", "uf|reajustable", "2 anos|720"]),
    # EEE: actividad (año en curso / año siguiente)
    "eee_pib_0": dict(frec=["MONTHLY"], palabras=["pib", "curso|actual|diciembre de este"], excluir=["no miner", "trimestre", "imacec"]),
    "eee_pib_1": dict(frec=["MONTHLY"], palabras=["pib", "siguiente|proximo|diciembre del proximo"], excluir=["no miner", "trimestre", "imacec"]),
    "eee_con_0": dict(frec=["MONTHLY"], palabras=["consumo", "curso|actual"]),
    "eee_con_1": dict(frec=["MONTHLY"], palabras=["consumo", "siguiente|proximo"]),
    "eee_fbcf_0": dict(frec=["MONTHLY"], palabras=["fbcf|formacion bruta|inversion", "curso|actual"]),
    "eee_fbcf_1": dict(frec=["MONTHLY"], palabras=["fbcf|formacion bruta|inversion", "siguiente|proximo"]),
    # Imacec no minero (índice original)
    "imacec_nm": dict(frec=["MONTHLY"], palabras=["imacec", "no miner"], excluir=["desestacionaliz", "variaci", "incidencia", "serie empalmada"]),
}


def normalizar(texto):
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


###############################################################################
# Cliente BDE                                                                 #
###############################################################################
class BDE:
    def __init__(self, user, password):
        self.user, self.password = user, password
        self.s = requests.Session()
        os.makedirs(CACHE, exist_ok=True)

    def _get(self, **params):
        params.update(user=self.user, pass_=self.password)
        params["pass"] = params.pop("pass_")
        for intento in range(4):
            try:
                r = self.s.get(BDE_URL, params=params, timeout=60)
                r.raise_for_status()
                js = r.json()
                if js.get("Codigo", 0) != 0:
                    raise RuntimeError(f"BDE respondió {js.get('Codigo')}: {js.get('Descripcion')}")
                return js
            except (requests.ConnectionError, requests.Timeout):
                if intento == 3:
                    raise
                time.sleep(2 ** (intento + 1))

    def catalogo(self, frec):
        ruta = os.path.join(CACHE, f"catalogo_{frec}.json")
        if os.path.exists(ruta) and time.time() - os.path.getmtime(ruta) < 7 * 86400:
            with open(ruta, encoding="utf-8") as f:
                return json.load(f)
        js = self._get(function="SearchSeries", frequency=frec)
        infos = js.get("SeriesInfos") or []
        with open(ruta, "w", encoding="utf-8") as f:
            json.dump(infos, f, ensure_ascii=False)
        return infos

    def serie(self, codigo, desde="1990-01-01"):
        js = self._get(function="GetSeries", timeseries=codigo, firstdate=desde,
                       lastdate=dt.date.today().isoformat())
        obs = (js.get("Series") or {}).get("Obs") or []
        df = pd.DataFrame(obs)
        if df.empty:
            return pd.Series(dtype=float, name=codigo)
        idx = pd.to_datetime(df["indexDateString"], format="%d-%m-%Y")
        val = pd.to_numeric(df["value"], errors="coerce")
        return pd.Series(val.values, index=idx, name=codigo).dropna().sort_index()


def coincide(titulo, palabras, excluir):
    t = normalizar(titulo)
    if any(normalizar(e) in t for e in excluir):
        return False
    return all(any(normalizar(alt) in t for alt in p.split("|")) for p in palabras)


def resolver(bde, clave, spec, log):
    frecs = spec.get("frec", ["MONTHLY"])
    if spec.get("codigo"):
        log.append(dict(clave=clave, codigo=spec["codigo"], titulo="(código fijo)", candidatos=1))
        return spec["codigo"]
    candidatos = []
    for frec in frecs:
        for info in bde.catalogo(frec):
            if coincide(info.get("spanishTitle", ""), spec["palabras"], spec.get("excluir", [])):
                candidatos.append(info)
        if candidatos:
            break
    if not candidatos:
        log.append(dict(clave=clave, codigo="", titulo="NO ENCONTRADA", candidatos=0))
        return None
    # Preferir la serie vigente (última observación más reciente) y el título más corto.
    def orden(i):
        ultima = pd.to_datetime(i.get("lastObservation"), dayfirst=True, errors="coerce")
        return (-(ultima.value if pd.notna(ultima) else 0), len(i.get("spanishTitle", "")))
    candidatos.sort(key=orden)
    elegido = candidatos[0]
    log.append(dict(clave=clave, codigo=elegido["seriesId"], titulo=elegido["spanishTitle"],
                    candidatos=len(candidatos),
                    alternativas=" || ".join(f'{c["seriesId"]}: {c["spanishTitle"]}' for c in candidatos[1:6])))
    return elegido["seriesId"]


def descargar_gscpi():
    r = requests.get(GSCPI_URL, timeout=60, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    return leer_gscpi(io.BytesIO(r.content))


def leer_gscpi(archivo):
    hojas = pd.read_excel(archivo, sheet_name=None, header=None)
    hoja = next((h for n, h in hojas.items() if "month" in n.lower()), next(iter(hojas.values())))
    fechas = pd.to_datetime(hoja.iloc[:, 0], errors="coerce")
    valores = pd.to_numeric(hoja.iloc[:, 1], errors="coerce")
    s = pd.Series(valores.values, index=fechas).dropna()
    s = s[s.index.notna()]
    s.index = s.index.to_period("M").to_timestamp()
    return s.sort_index()


###############################################################################
# Gráficos                                                                    #
###############################################################################
def fmt_mes(sep="-", anio2=True):
    def f(x, _):
        d = mdates.num2date(x)
        a = f"{d.year % 100:02d}" if anio2 else str(d.year)
        return f"{MESES[d.month - 1]}{sep}{a}"
    return mticker.FuncFormatter(f)


def fmt_miles(x, _):
    return f"{x:,.0f}".replace(",", ".")


def fmt_dec(dec=None):
    if dec is None:
        return mticker.FuncFormatter(lambda x, _: f"{x:g}".replace(".", ","))
    return mticker.FuncFormatter(lambda x, _: f"{x:.{dec}f}".replace(".", ","))


def nuevo_grafico(ancho, alto):
    fig, ax = plt.subplots(figsize=(ancho, alto), dpi=200)
    fig.subplots_adjust(left=0.09, right=0.95, top=0.97, bottom=0.12)
    ax.tick_params(length=3, color=GRIS_LINEA)
    return fig, ax


def etiqueta(ax, texto, x=0.04, y=0.95, ha="left"):
    ax.text(x, y, texto, transform=ax.transAxes, va="top", ha=ha, color="black", fontsize=8.5)


def linea_50(ax):
    ax.axhline(50, color=GRIS, lw=0.8, ls=(0, (3, 3)))


def eje_mensual(ax, intervalo, formato, rotar=False, mes_inicial=1):
    meses = sorted({(mes_inicial - 1 + i) % 12 + 1 for i in range(0, 12, intervalo)})
    ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=meses))
    ax.xaxis.set_major_formatter(formato)
    if rotar:
        plt.setp(ax.get_xticklabels(), rotation=90)


def recortar(s, desde):
    return s[s.index >= pd.Timestamp(desde)]


def guardar(fig, nombre):
    ruta = os.path.join(CACHE, "graficos", nombre + ".png")
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    fig.savefig(ruta, dpi=200, transparent=False, facecolor="white")
    plt.close(fig)
    return ruta


def g_lineas(nombre, series, ancho, alto, desde, ylim=None, yticks=None, etiqueta_txt=None,
             fmt_x=None, intervalo=12, rotar=False, linea50=False, leyenda="lower center",
             ncol=1, dec=None, lw=1.4, inicio_eje_anual=False, pos_etiqueta=0.95, ancla_leyenda=None):
    fig, ax = nuevo_grafico(ancho, alto)
    fin = pd.Timestamp(desde)
    for etiq, s, color in series:
        if s is None or s.empty:
            continue
        s = recortar(s, desde)
        ax.plot(s.index, s.values, color=color, lw=lw, label=etiq)
        fin = max(fin, s.index.max())
    if linea50:
        linea_50(ax)
    if ylim:
        ax.set_ylim(*ylim)
    if yticks is not None:
        ax.set_yticks(yticks)
    ax.yaxis.set_major_formatter(fmt_dec(dec))
    if inicio_eje_anual:
        ax.xaxis.set_major_locator(mdates.YearLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
        if rotar:
            plt.setp(ax.get_xticklabels(), rotation=90)
    else:
        eje_mensual(ax, intervalo, fmt_x or fmt_mes(), rotar, pd.Timestamp(desde).month)
    ax.set_xlim(pd.Timestamp(desde), fin + pd.Timedelta(days=15))
    if etiqueta_txt:
        etiqueta(ax, etiqueta_txt, y=pos_etiqueta)
    if leyenda:
        ax.legend(loc=leyenda, ncol=ncol, handlelength=1.6, bbox_to_anchor=ancla_leyenda)
    return guardar(fig, nombre)


def anual_desde_curso_y_siguiente(curso, siguiente, anio):
    """Expectativa para el año `anio`: 'año siguiente' durante anio-1 y 'año en curso' durante anio."""
    partes = []
    if siguiente is not None:
        partes.append(siguiente[siguiente.index.year == anio - 1])
    if curso is not None:
        partes.append(curso[curso.index.year == anio])
    partes = [p for p in partes if not p.empty]
    return pd.concat(partes).sort_index() if partes else pd.Series(dtype=float)


###############################################################################
# Presentación                                                                #
###############################################################################
class Deck:
    W, H = Inches(13.333), Inches(7.5)

    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = self.W, self.H
        self.blank = self.prs.slide_layouts[6]
        self.n = 1

    def _texto(self, slide, x, y, w, h, texto, tam, color="000000", negrita=False,
               alinear=PP_ALIGN.LEFT, fuente=FUENTE):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, linea in enumerate(texto.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = alinear
            r = p.add_run()
            r.text = linea
            r.font.size, r.font.bold, r.font.name = Pt(tam), negrita, fuente
            r.font.color.rgb = RGBColor.from_string(color)
        return tb

    def _linea(self, slide, x1, y, x2, color="D9D9D9", grosor=0.5):
        ln = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y))
        ln.line.color.rgb = RGBColor.from_string(color)
        ln.line.width = Pt(grosor)

    def portada(self, fecha_txt):
        s = self.prs.slides.add_slide(self.blank)
        s.shapes.add_picture(os.path.join(ASSETS, "portada.jpg"), 0, 0, self.W, self.H)
        s.shapes.add_picture(os.path.join(ASSETS, "logo_portada.png"),
                             Inches(11.29), Inches(0.12), Inches(1.6), Inches(2.01))
        self._texto(s, 9.42, 3.25, 3.5, 1.4, "Estudios\nEconómicos", 32, "EC0000", fuente="Calibri Light")
        self._linea(s, 9.37, 4.64, 12.27, color="1F4E79", grosor=1)
        self._texto(s, 9.42, 4.72, 3.0, 0.3, fecha_txt, 12, "EC0000")

    def lamina(self, titulo, graficos, fuentes, layout):
        """graficos: lista de (subtitulo, ruta_png). fuentes: texto o lista por columna."""
        self.n += 1
        s = self.prs.slides.add_slide(self.blank)
        self._texto(s, 0.5, 0.2, 10.5, 0.5, titulo, 20, "EC0000", negrita=True)
        s.shapes.add_picture(os.path.join(ASSETS, "logo_llama.png"),
                             Inches(12.62), Inches(0), Inches(0.46), Inches(0.6))
        self._linea(s, 0.5, 0.8, 12.35)
        self._linea(s, 0.5, 6.9, 12.8)
        for (sub, png), (x, w) in zip(graficos, layout):
            self._texto(s, x, 1.3, w, 0.6, sub, 11)
            s.shapes.add_picture(png, Inches(x), Inches(2.0), Inches(w))
        if isinstance(fuentes, str):
            fuentes = [fuentes]
        for f, (x, _) in zip(fuentes, layout):
            if f:
                self._texto(s, x, 7.02, 5, 0.2, f"Fuente: {f}", 7)
        self._texto(s, 12.2, 7.05, 0.6, 0.3, str(self.n), 14, "7F7F7F", negrita=True, alinear=PP_ALIGN.RIGHT)

    def guardar(self, ruta):
        self.prs.save(ruta)


DOS = [(0.5, 5.85), (6.9, 5.85)]
TRES = [(0.5, 3.9), (4.6, 3.9), (8.9, 3.9)]
UNO = [(0.5, 11.9)]
ALTO = 4.85  # alto de los gráficos (pulgadas)


###############################################################################
# Principal                                                                   #
###############################################################################
def construir(datos, hoy):
    """datos: dict clave -> pd.Series. Devuelve ruta del PPTX."""
    d = datos.get
    a1 = hoy.year - 1  # inicio de ventanas cortas (enero del año anterior)
    deck = Deck()
    deck.portada(f"{MESES_LARGO[hoy.month - 1]} de {hoy.year}")
    BCCH = "Banco Central de Chile"
    w2, w3, w1 = DOS[0][1], TRES[0][1], UNO[0][1]

    # 2. IMCE
    kw = dict(desde="2018-01-01", ylim=(0, 100), yticks=range(0, 101, 10), etiqueta_txt="(Índice)",
              intervalo=12, linea50=True, leyenda="lower right", ncol=3)
    imce = [("Comercio", AZUL_CLARO), ("Construcción", VERDE), ("Industria", ROJO)]
    deck.lamina("IMCE", [
        ("IMCE: Indicador de Precios", g_lineas("imce_precios", [(n, d(k), c) for (n, c), k in zip(imce, ["imce_pre_com", "imce_pre_con", "imce_pre_ind"])], w2, ALTO, **kw)),
        ("IMCE: Indicador de Costos", g_lineas("imce_costos", [(n, d(k), c) for (n, c), k in zip(imce, ["imce_cos_com", "imce_cos_con", "imce_cos_ind"])], w2, ALTO, **kw)),
    ], BCCH, DOS)

    # 3. Expectativas de inflación
    kw = dict(desde="2015-01-01", ylim=(1, 9), yticks=range(1, 10), etiqueta_txt="(Var. anual, %)",
              leyenda="upper left", ancla_leyenda=(0.08, 0.83), ncol=2, inicio_eje_anual=True)
    deck.lamina("Expectativas de Inflación", [
        ("Mediana EEE 1 y 2 años\n(Var. anual, %)", g_lineas("eee_inflacion", [("EEE 1 AÑO", d("eee_1"), AZUL), ("EEE 2 AÑOS", d("eee_2"), ROJO_OSC)], w2, ALTO, **kw)),
        ("Mediana EOF 1 y 2 años\n(Var. anual, %)", g_lineas("eof_inflacion", [("EOF 1 AÑO", d("eof_1"), AZUL), ("EOF 2 AÑOS", d("eof_2"), ROJO_OSC)], w2, ALTO, **kw)),
    ], BCCH, DOS)

    # 4. Compensación inflacionaria
    kw = dict(desde="2022-01-01", intervalo=6, fmt_x=fmt_mes(".", True), lw=1.2)
    deck.lamina("Compensación Inflacionaria", [
        ("Compensaciones inflacionarias\n(Var. Anual, %)", g_lineas("compensaciones", [("1 año", d("comp_1"), ROJO), ("2 años", d("comp_2"), PETROLEO)], w2, ALTO, etiqueta_txt="(Var. anual, %)", leyenda="upper right", ncol=2, **kw)),
        ("Compensación inflacionaria 1 en 1: Swap\n(%)", g_lineas("comp_1en1", [("1 en 1", d("comp_1en1_swap"), AZUL)], w2, ALTO, etiqueta_txt="(%)", leyenda=None, **kw)),
    ], BCCH, DOS)

    # 5. EDEP / EPN
    kw = dict(desde="2021-11-01", intervalo=6, fmt_x=fmt_mes(".", True), rotar=True, lw=1.6, pos_etiqueta=0.1)
    costos = [("Costos laborales", ROJO_OSC), ("Precios de insumos", DORADO),
              ("Costo de maquinarias o\ntécnología", VERDE_SUAVE), ("Costos fijos", CELESTE)]
    deck.lamina("EDEP", [
        ("EDEP: Evolución costos de la empresa\n(índice de difusión)",
         g_lineas("edep_costos", [(n, d(k), c) for (n, c), k in zip(costos, ["edep_cos_lab", "edep_cos_ins", "edep_cos_maq", "edep_cos_fij"])],
                  w3, ALTO, linea50=True, etiqueta_txt="(Índice de difusión)", leyenda="upper right", **kw)),
        ("EDEP: Influencia sobre precio de venta\n(índice de difusión)",
         g_lineas("edep_influencia", [("Niveles de venta", d("edep_inf_ven"), "#000000"), ("Precio del dólar", d("edep_inf_dol"), NARANJO),
                                      ("Márgenes de ganancias", d("edep_inf_mar"), MORADO), ("Disponibilidad de insumos", d("edep_inf_dis"), CELESTE),
                                      ("Costos", d("edep_inf_cos"), GRIS)],
                  w3, ALTO, linea50=True, etiqueta_txt="(Índice de difusión)", leyenda="lower right", **{**kw, "pos_etiqueta": 0.95})),
        ("EPN: Evolución de los costos\n(índice de difusión, percepción evolución\núltimos 3 meses)",
         g_lineas("epn_costos", [(n, d(k), c) for (n, c), k in zip([("Costos laborales", "#000000"), ("Precios de insumos", DORADO), costos[2], ("Costos fijos", AZUL_CLARO)], ["epn_cos_lab", "epn_cos_ins", "epn_cos_maq", "epn_cos_fij"])],
                  w3, ALTO, linea50=True, etiqueta_txt="(Índice de difusión)", leyenda="upper right", **kw)),
    ], BCCH, TRES)

    # 6. EDEP expectativas y GSCPI
    deck.lamina("IMCE, EDEP y Costos de transporte", [
        ("EDEP: Expectativas de Inflación\n(Var. Anual, %)",
         g_lineas("edep_expectativas", [("EDEP a 1 año", d("edep_inf_1"), PETROLEO), ("EDEP a 2 años", d("edep_inf_2"), NARANJO)], w2, ALTO,
                  desde="2021-11-01", intervalo=6, fmt_x=fmt_mes(".", True), ylim=(0, 12), yticks=range(0, 13, 2),
                  etiqueta_txt="(Var. anual, %)", leyenda="lower center", ncol=2)),
        ("Costos de transporte\n(desviación estándar del valor promedio)",
         g_lineas("gscpi", [("GSCPI", d("gscpi"), AZUL)], w2, ALTO, desde="1998-01-01", inicio_eje_anual=True, rotar=True,
                  etiqueta_txt="(Desviación estandar del valor\npromedio)", leyenda=None, lw=1.0)),
    ], [BCCH, "Reserva Federal de Nueva York"], DOS)

    # 7. TCN / TCM (doble eje)
    fig, ax = nuevo_grafico(w1, ALTO)
    fig.subplots_adjust(left=0.05, right=0.95)
    tcn, tcm = recortar(d("tcn", pd.Series(dtype=float)), "2018-01-01"), recortar(d("tcm", pd.Series(dtype=float)), "2018-01-01")
    ax2 = ax.twinx()
    ax2.spines["right"].set_visible(True)
    ax2.spines["top"].set_visible(False)
    ax.plot(tcn.index, tcn.values, color=ROJO_OSC, lw=1.2, label="TCN")
    ax2.plot(tcm.index, tcm.values, color=AZUL, lw=1.2, label="TCM (DERECHA)")
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(fmt_miles))
    if not tcn.empty:
        ax.set_xlim(tcn.index.min(), max(tcn.index.max(), tcm.index.max() if not tcm.empty else tcn.index.max()))
    ax.xaxis.set_major_locator(mdates.YearLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    ax2.tick_params(colors=GRIS, length=3)
    etiqueta(ax, "(Pesos chilenos)")
    etiqueta(ax2, "(Índice)", x=0.96, ha="right")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc="lower center", bbox_to_anchor=(0.6, 0.12), ncol=2)
    deck.lamina("TCN TCM", [("Tipo de cambio nominal spread Tipo de cambio multilateral\n(Pesos chilenos) (Índice)", guardar(fig, "tcn_tcm"))], BCCH, UNO)

    # 8. SPC
    kw = dict(desde=f"{a1}-01-01", intervalo=3, fmt_x=fmt_mes(".", True), dec=1, leyenda="upper right", ncol=2, lw=1.2)
    deck.lamina("Tasas SPC", [
        ("SPC pesos: 1 y 2 años\n(%)", g_lineas("spc_pesos", [("SPC pesos 2 años", d("spc_clp_2"), AZUL), ("SPC pesos 1 año", d("spc_clp_1"), ROJO)], w2, ALTO, etiqueta_txt="(%)", **kw)),
        ("SPC UF 1 y 2 años\n(%)", g_lineas("spc_uf", [("SPC UF 2 años", d("spc_uf_2"), AZUL), ("SPC UF 1 año", d("spc_uf_1"), ROJO)], w2, ALTO, etiqueta_txt="(%)", **kw)),
    ], BCCH, DOS)

    # 9. EEE actividad: año en curso y año siguiente, armados por año calendario
    a0 = hoy.year
    kw = dict(desde=f"{a1}-01-01", intervalo=3, fmt_x=fmt_mes(".", True), dec=1, etiqueta_txt="(Var. anual, %)",
              leyenda="lower right", ncol=1, lw=1.6)
    def par(cur, sig):
        return [(str(a0 + 1), anual_desde_curso_y_siguiente(d(cur), d(sig), a0 + 1), PETROLEO),
                (str(a0), anual_desde_curso_y_siguiente(d(cur), d(sig), a0), DORADO)]
    deck.lamina("EEE", [
        ("Expectativas de PIB\n(Var. Anual, %)", g_lineas("eee_pib", par("eee_pib_0", "eee_pib_1"), w3, ALTO, **kw)),
        ("Expectativas de Consumo de los hogares\n(Var. Anual, %)", g_lineas("eee_consumo", par("eee_con_0", "eee_con_1"), w3, ALTO, **kw)),
        ("Expectativas FBCF\n(Var. Anual, %)", g_lineas("eee_fbcf", par("eee_fbcf_0", "eee_fbcf_1"), w3, ALTO, **kw)),
    ], BCCH, TRES)

    # 10. IMCE expectativas de inversión vs Imacec no minero (var. trimestral a/a)
    fig, ax = nuevo_grafico(w1, ALTO)
    fig.subplots_adjust(left=0.05, right=0.95, bottom=0.14)
    imacec = d("imacec_nm", pd.Series(dtype=float))
    if not imacec.empty:
        trim = imacec.rolling(3).mean()
        imacec = (trim / trim.shift(12) - 1) * 100
    imacec = recortar(imacec.dropna(), "2010-01-01")
    inv = recortar(d("imce_inv", pd.Series(dtype=float)), "2010-01-01")
    ax2 = ax.twinx()
    ax2.spines["right"].set_visible(True)
    ax2.spines["top"].set_visible(False)
    ax.plot(imacec.index, imacec.values, color=DORADO, lw=1.8, label="Imacec No Minero")
    ax2.plot(inv.index, inv.values, color=AZUL, lw=1.8, label="IMCE expectativas de inversión (derecha)")
    fin = max([s.index.max() for s in (imacec, inv) if not s.empty], default=pd.Timestamp("2010-01-01"))
    ax.set_xlim(pd.Timestamp("2010-01-01"), fin + pd.Timedelta(days=30))
    ax.xaxis.set_major_locator(mdates.MonthLocator(bymonth=[1, 7]))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    plt.setp(ax.get_xticklabels(), rotation=90)
    ax2.tick_params(colors=GRIS, length=3)
    etiqueta(ax, "(Var. Trimestral a/a)", x=0.12)
    etiqueta(ax2, "(Índice)", x=0.93, ha="right")
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc="lower left", bbox_to_anchor=(0.06, 0.18), ncol=2)
    deck.lamina("Imce expectativas v/s Imacec no minero",
                [("Imce expectativas v/s Imacec no minero\n(Índice) (Var. Trimestral a/a)", guardar(fig, "imce_imacec"))], BCCH, UNO)

    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, f"{hoy:%Y%m%d}_Chartpack_Expectativas.pptx")
    deck.guardar(ruta)
    return ruta


def a_pdf(ruta_pptx):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice no está instalado: se omite el PDF.")
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", SALIDA, ruta_pptx],
                   check=True, capture_output=True)
    return os.path.splitext(ruta_pptx)[0] + ".pdf"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--user", default=os.environ.get("BCCH_USER"))
    ap.add_argument("--password", default=os.environ.get("BCCH_PASS"))
    ap.add_argument("--buscar", help="busca texto en el catálogo de la BDE y termina")
    ap.add_argument("--gscpi", help="ruta local a gscpi_data.xlsx (si no se puede descargar)")
    args = ap.parse_args()
    if not args.user or not args.password:
        sys.exit("Faltan credenciales: defina BCCH_USER y BCCH_PASS (o --user/--password).")

    configurar_fuente()
    bde = BDE(args.user, args.password)

    if args.buscar:
        for frec in ["DAILY", "MONTHLY", "QUARTERLY", "ANNUAL"]:
            for i in bde.catalogo(frec):
                if coincide(i.get("spanishTitle", ""), args.buscar.split(), []):
                    print(f'{frec:9s} {i["seriesId"]:40s} {i.get("lastObservation", ""):12s} {i["spanishTitle"]}')
        return

    log, datos = [], {}
    for clave, spec in SERIES.items():
        codigo = resolver(bde, clave, spec, log)
        if codigo:
            try:
                datos[clave] = bde.serie(codigo, desde="1997-01-01")
                log[-1]["ultimo_dato"] = datos[clave].index.max().date() if not datos[clave].empty else ""
            except Exception as e:  # una serie con problemas no detiene el chartpack
                log[-1]["titulo"] += f" | ERROR: {e}"
    try:
        datos["gscpi"] = leer_gscpi(args.gscpi) if args.gscpi else descargar_gscpi()
        log.append(dict(clave="gscpi", codigo="NY Fed GSCPI", titulo="Global Supply Chain Pressure Index",
                        candidatos=1, ultimo_dato=datos["gscpi"].index.max().date()))
    except Exception as e:
        log.append(dict(clave="gscpi", codigo="", titulo=f"ERROR: {e}", candidatos=0))

    os.makedirs(SALIDA, exist_ok=True)
    pd.DataFrame(log).to_csv(os.path.join(SALIDA, "series_resueltas.csv"), index=False, encoding="utf-8-sig")
    faltantes = [r["clave"] for r in log if not r.get("codigo")]
    if faltantes:
        print("ATENCIÓN, series sin resolver (fije `codigo` en SERIES):", ", ".join(faltantes))

    pptx = construir(datos, dt.date.today())
    print("PPTX:", pptx)
    pdf = a_pdf(pptx)
    if pdf:
        print("PDF: ", pdf)


if __name__ == "__main__":
    main()
